#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
llm-privacy-gateway · 本地文字提取模块 (local_extract.py)

文件/图片正文在本地完成解析与文字提取，文件本体不出域：
  - .txt/.md/.csv/.json/.log  直接读文本
  - .pdf                       pypdf 提取
  - .docx                      python-docx 提取（段落+表格）
  - .xlsx                      openpyxl 提取（全部单元格）
  - 图片(.png/.jpg/.jpeg/.bmp/.webp/.tiff)
                               OCR 提取（优先 paddleocr，其次 pytesseract；
                               未安装引擎时给出明确安装指引）

用法:
  python local_extract.py --file path [--max-chars 20000]
"""
import argparse
import os
import sys

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tiff"}
TEXT_EXTS = {".txt", ".md", ".csv", ".json", ".log", ".xml", ".html"}
SUPPORTED_EXTS = (TEXT_EXTS | IMAGE_EXTS | {".pdf", ".docx", ".xlsx"})

_DEP_HINT = "请先执行  python scripts/setup_deps.py  补齐依赖"


def extract_text(path, max_chars=None):
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTS:
        raise ValueError(
            "不支持的文件类型: %s（支持：txt/md/csv/json 等文本、PDF、Word(.docx)、"
            "Excel(.xlsx)、常见图片；.doc/.xls 旧版格式请先另存为新格式）" % ext)
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    if ext in TEXT_EXTS:
        text = _read_text_file(path)
    elif ext == ".pdf":
        text = _extract_pdf(path)
    elif ext == ".docx":
        text = _extract_docx(path)
    elif ext == ".xlsx":
        text = _extract_xlsx(path)
    elif ext in IMAGE_EXTS:
        text = _extract_image(path)
    else:
        raise ValueError("不支持的文件类型: %s" % ext)
    if max_chars and len(text) > max_chars:
        text = text[:max_chars] + "\n...[已截断]"
    return text


def _read_text_file(path):
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    # 兜底：忽略错误解码
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("缺少 PDF 解析依赖（%s）。" % _DEP_HINT)
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:  # 单页失败不中断整体提取
            parts.append("[第 %d 页提取失败: %s]" % (len(parts) + 1, exc))
    return "\n".join(p for p in parts if p)


def _extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("缺少 Word 解析依赖（%s）。" % _DEP_HINT)
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("缺少 Excel 解析依赖（%s）。" % _DEP_HINT)
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append("== Sheet: %s ==" % ws.title)
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                parts.append(" | ".join(vals))
    wb.close()
    return "\n".join(parts)


def _extract_image(path):
    """OCR：优先 paddleocr，其次 pytesseract；两者都缺时给出明确指引。"""
    try:
        from paddleocr import PaddleOCR
        ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        result = ocr.ocr(path, cls=True)
        lines = []
        for page in result or []:
            for item in page or []:
                lines.append(item[1][0])
        return "\n".join(lines)
    except ImportError:
        pass
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path), lang="chi_sim+eng")
    except ImportError:
        pass
    except Exception as exc:
        raise RuntimeError("OCR 引擎调用失败：%s" % exc)
    raise RuntimeError(
        "未安装 OCR 引擎，无法提取图片文字。请任选其一：\n"
        "  1) python -m pip install paddleocr paddlepaddle   （推荐，中文效果好）\n"
        "  2) python -m pip install pytesseract ，并安装 Tesseract OCR 后设置语言包 chi_sim")


def main():
    ap = argparse.ArgumentParser(description="llm-privacy-gateway 本地文字提取")
    ap.add_argument("--file", required=True, help="本地文件路径")
    ap.add_argument("--max-chars", type=int, default=None)
    args = ap.parse_args()
    try:
        text = extract_text(args.file, max_chars=args.max_chars)
    except Exception as exc:
        print("提取失败：%s" % exc, file=sys.stderr)
        sys.exit(1)
    print("已从本地提取 %d 字符（文件本体未上传）。" % len(text), file=sys.stderr)
    print(text)


if __name__ == "__main__":
    main()
